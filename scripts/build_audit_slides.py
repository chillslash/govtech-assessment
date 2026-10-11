"""Build a 7-slide editable GovTech assessment proposal from frozen metrics.
Requires python-pptx (available in the local ignored .deps directory).
"""
from pathlib import Path
import json,sys,zipfile
sys.path.insert(0,str(Path(__file__).resolve().parents[1]/'.deps'))
from pptx import Presentation
from pptx.util import Inches,Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN,MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

ROOT=Path(__file__).resolve().parents[1]
summary=json.loads((ROOT/'data/evaluations/anao-full-evaluation-summary.json').read_text(encoding='utf8'))
a=summary['totals']
assert (a['screened'],a['nominated'],a['verified'],a['ref_total'])==(189,102,286,58)

W=13.333;H=7.5
prs=Presentation()
prs.slide_width=Inches(W);prs.slide_height=Inches(H)
blank=prs.slide_layouts[6]
def rgb(h):return RGBColor.from_string(h.strip('#'))
NAVY='152340'; DARK='1B2943'; SLATE='526278'; LIGHT='F5F7FB'; WHITE='FFFFFF'
BLUE='4467E8'; TEAL='14A996'; PURPLE='775CD6'; BORDER='E3E8F0'; AMBER='D99B35'
FONT='Aptos'
def shape(s,x,y,w,h,fill=WHITE,line=None,round=False):
    sh=s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if round else MSO_SHAPE.RECTANGLE,
                          Inches(x),Inches(y),Inches(w),Inches(h))
    sh.fill.solid();sh.fill.fore_color.rgb=rgb(fill)
    if line:sh.line.color.rgb=rgb(line);sh.line.width=Pt(1)
    else:sh.line.fill.background()
    return sh
def text(s,t,x,y,w,h,size=18,color=DARK,bold=False,align=None,margin=0.01,
         valign=MSO_ANCHOR.TOP):
    tx=s.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h))
    tf=tx.text_frame;tf.clear();tf.word_wrap=True
    tf.margin_left=Inches(margin);tf.margin_right=Inches(margin)
    tf.margin_top=Inches(margin);tf.margin_bottom=Inches(margin)
    tf.vertical_anchor=valign
    lines=str(t).split('\n')
    for i,line in enumerate(lines):
        p=tf.paragraphs[0] if i==0 else tf.add_paragraph()
        p.text=line;p.font.name=FONT;p.font.size=Pt(size);p.font.bold=bold
        p.font.color.rgb=rgb(color);p.space_before=Pt(0);p.space_after=Pt(0)
        if align:p.alignment=align
    return tx
def base(title,section,no,subtitle=None):
    s=prs.slides.add_slide(blank)
    shape(s,0,0,W,H,LIGHT)
    text(s,section.upper(),.65,.34,10,.26,10,BLUE,True)
    text(s,title,.65,.78,12.15,.61,30,NAVY,True)
    if subtitle:text(s,subtitle,.68,1.52,11.95,.52,14,SLATE)
    shape(s,.65,7.06,12.05,.014,BORDER)
    text(s,'EVIDENCE-FIRST AUDIT ASSURANCE  /  PUBLIC ANAO REPORTS',.67,7.13,9,.18,8,SLATE)
    text(s,str(no).zfill(2)+' / 07',11.87,7.11,.8,.18,9,SLATE,True,PP_ALIGN.RIGHT)
    return s
def label(s,t,x,y,w,color=SLATE):
    text(s,t.upper(),x,y,w,.22,10,color,True)
def card(s,x,y,w,h,fill=WHITE):
    return shape(s,x,y,w,h,fill,BORDER,True)
def tiny_rule(s,x,y,w,col=BLUE):shape(s,x,y,w,.045,col)
def badge(s,t,x,y,w,col=BLUE,fill='ECF0FF'):
    shape(s,x,y,w,.38,fill,None,True)
    text(s,t,x+.1,y+.075,w-.2,.22,11,col,True)

# 1 Cover
s=prs.slides.add_slide(blank)
shape(s,0,0,W,H,NAVY)
shape(s,8.57,-.45,4.2,8.4,'223A66')
shape(s,9.25,1.15,2.25,2.25,'31518B',None,True)
shape(s,10.03,2.15,2.25,2.25,'416DE0',None,True)
shape(s,8.5,4.85,3.15,.08,TEAL)
badge(s,'GOVTECH TECHNICAL ASSESSMENT',.7,.66,3.75,'D6E6FF','294166')
text(s,'From audit reports\nto defensible insights',.72,1.48,8.38,2.14,40,WHITE,True)
text(s,'An evidence-first LLM workflow for discovering, verifying and comparing IT-control weaknesses across reports.',.74,4.03,7.53,1.1,19,'DAE3F4')
tiny_rule(s,.74,5.52,.86,TEAL)
text(s,'Five published audits   •   Working prototype   •   Independent source review',.74,5.78,8.58,.38,14,WHITE)
text(s,'11 OCTOBER 2026     |     PROPOSAL & EVALUATION',.74,7.05,8.1,.24,10,'C6D2E8',True)

# 2 Problem
s=base('The goal is not another summariser','01  /  The audit workflow',2,
       'Auditors need a defensible finding register before they can compare risks across reports.')
card(s,.68,2.22,5.74,3.92)
label(s,'Typical report review',.94,2.48,4.8,'A15D40')
for i,(hd,body) in enumerate([
 ('Scattered evidence','Relevant IT-control weaknesses appear under different headings and layouts.'),
 ('Non-atomic observations','One passage can describe several independent failures.'),
 ('Unverifiable comparisons','Similar language across reports does not prove identical underlying problems.')]):
    yy=2.92+i*1.02
    shape(s,.96,yy+.04,.09,.09,'CB8254',None,True)
    text(s,hd,1.2,yy,4.95,.29,17,DARK,True)
    text(s,body,1.2,yy+.35,4.84,.6,12,SLATE)
card(s,6.78,2.22,5.88,3.92)
label(s,'Proposed reviewer experience',7.08,2.48,4.85,TEAL)
for i,(hd,body) in enumerate([
 ('Extract explicit weaknesses','Structured, separately remediable records with original citations.'),
 ('Validate what software can','Check page, quote, schema and original paragraph/table locator.'),
 ('Compare with human control','Surface candidate themes with traceable IDs; auditors approve interpretations.')]):
    yy=2.92+i*1.02
    shape(s,7.09,yy+.04,.09,.09,TEAL,None,True)
    text(s,hd,7.33,yy,4.8,.29,17,DARK,True)
    text(s,body,7.33,yy+.35,4.88,.6,12,SLATE)
text(s,'DESIGN PRINCIPLE  A cross-report theme is only useful if its source records can be reopened and challenged.',.82,6.38,11.72,.42,14,NAVY,True)

# 3 Evidence
s=base('One quote becomes a reviewable record','02  /  Concrete example',3,
       'A real observation recovered from a page excluded by the original heuristic shortlist.')
card(s,.68,2.19,5.95,3.78)
label(s,'Original report evidence',.99,2.53,4.75,BLUE)
text(s,'“logging has not been configured for Australia Post’s desktops”',.99,3.13,5.28,1.34,26,NAVY,True)
shape(s,.98,4.73,4.99,.02,BORDER)
text(s,'ANAO 2019/20 corporate cyber-resilience audit\nPhysical PDF page 32',1.0,4.93,5.1,.69,13,SLATE)
card(s,6.84,2.19,5.82,3.78)
label(s,'Structured finding (proposed)',7.12,2.53,4.87,TEAL)
items=[('ENTITY','Australia Post'),('CONTROL AREA','Logging / monitoring'),
       ('ATOMIC WEAKNESS','Desktop event logging not configured'),
       ('EVIDENCE ANCHOR','Source PDF page 32 + verbatim passage')]
for i,(k,v) in enumerate(items):
    yy=2.96+i*.64
    text(s,k,7.12,yy,1.72,.24,10,SLATE,True)
    text(s,v,8.94,yy,3.37,.51,14,NAVY,True if i<2 else False)
badge(s,'AUDITOR REVIEW REQUIRED',7.13,5.54,2.84,TEAL,'E6F7F3')
text(s,'Source verification proves the quotation exists; it does not by itself prove the model’s interpretation or agency attribution.',.84,6.24,11.52,.57,14,SLATE)

# 4 Architecture
s=base('Simple architecture, explicit responsibilities','03  /  System design',4,
       'Parsing and rule-checking remain deterministic; semantic decisions are separated and reviewable.')
stages=[
 ('01','PARSE','PDF → page text','CODE',BLUE),
 ('02','SCREEN','LLM examines each findings page','LLM',PURPLE),
 ('03','EXTRACT','Atomic IT-control weaknesses','LLM',PURPLE),
 ('04','VERIFY','Quote + page + locator checks','CODE',BLUE),
 ('05','COMPARE','Category counts + theme proposals','CODE + LLM',PURPLE),
 ('06','APPROVE','Auditor reviews accepted records','HUMAN',TEAL)]
for i,(num,head,body,owner,col) in enumerate(stages):
    x=.64+i*2.05
    card(s,x,2.40,1.83,2.68)
    shape(s,x,2.40,1.83,.08,col)
    text(s,num,x+.16,2.67,.69,.37,20,col,True)
    text(s,head,x+.16,3.13,1.62,.35,16,NAVY,True)
    text(s,body,x+.16,3.64,1.53,.84,12,SLATE)
    label(s,owner,x+.16,4.72,1.55,col)
    if i<5:text(s,'→',x+1.84,3.43,.22,.35,18,SLATE,True,PP_ALIGN.CENTER)
shape(s,.72,5.46,11.9,.87,'EAEFFD',None,True)
text(s,'Expanded test: all 189 substantive findings pages were screened. Detailed extraction continued only for pages nominated by discovery.',.99,5.68,11.27,.49,15,NAVY,True)

# 5 Tradeoffs
s=base('Three deliberate design trade-offs','04  /  Engineering decisions',5,
       'The constraints make the MVP auditable—but do not eliminate semantic errors.')
choices=[
 ('COVERAGE','LLM screens every scoped findings page; no arbitrary top-k shortlist.',
  'Discovery may still overlook a weakness or reject a relevant quote.',BLUE,'More complete inputs ≠ full recall'),
 ('GROUNDING','Every retained record has a source quote, physical page and checked locator.',
  'A real quote can be attached to the wrong entity or an overbroad claim.',PURPLE,'Valid citation ≠ true interpretation'),
 ('SYNTHESIS','Software groups candidates; LLM proposes semantic links backed by exact IDs.',
  'Matching topics alone do not establish a recurring control failure.',TEAL,'Theme proposal ≠ audit conclusion')]
for i,(head,positive,risk,col,principle) in enumerate(choices):
    x=.69+i*4.15
    card(s,x,2.28,3.91,3.83)
    shape(s,x,2.28,3.91,.075,col)
    label(s,head,x+.24,2.66,3.2,col)
    text(s,positive,x+.24,3.08,3.40,.98,16,NAVY,True)
    shape(s,x+.24,4.2,3.4,.017,BORDER)
    text(s,risk,x+.24,4.44,3.36,.81,13,SLATE)
    text(s,principle,x+.24,5.52,3.32,.42,11,col,True)
text(s,'Ingestion of confidential audit reports would require separate security, access-control, retention and vendor approval.',.82,6.48,11.70,.37,13,SLATE)

# 6 Results
s=base('Broader coverage, but accuracy is still unproven','05  /  Measured evaluation',6)
for i,(big,desc) in enumerate([
 (str(a['screened']),'Findings pages screened'),
 (str(a['nominated']),'Pages nominated for extraction'),
 (str(a['verified']),'Source-verified model records')]):
    x=.7+i*4.15
    card(s,x,1.74,3.9,1.40)
    text(s,big,x+.24,1.98,1.43,.67,35,BLUE,True)
    text(s,desc,x+1.51,2.05,2.08,.73,13,SLATE,True)
card(s,.7,3.39,7.75,2.54)
label(s,'Frozen 58-concept reference set: expanded model',.98,3.68,6.65)
segments=[(43,BLUE),(1,PURPLE),(2,AMBER),(12,'BBC4D2')]
x=.99
for n,col in segments:
    w=6.99*n/58
    shape(s,x,4.23,w,.36,col)
    x+=w
text(s,'43 clear     1 merged     2 uncertain     12 absent',.99,4.77,6.8,.4,15,NAVY,True)
text(s,'AI-assisted comparison on the original 14 labelled pages; not auditor-certified precision or report-level recall.',.99,5.27,7.02,.47,11,SLATE)
card(s,8.65,3.39,3.95,2.54)
label(s,'Previously excluded pages',8.93,3.68,3.35,TEAL)
text(s,'4 detected  ·  2 partial',8.93,4.23,3.33,.47,20,NAVY,True)
text(s,'1 missed among 7 positive\nholdout pages',8.93,4.80,3.34,.64,14,SLATE)
text(s,'Earlier bounded MVP: 50 source-verified records from 14 extracted pages. More output is not evidence of higher precision.',.83,6.26,11.52,.50,13,SLATE)

# 7 Pilot roadmap and themes
s=base('Theme proposals, then auditor-led validation','06  /  Responsible rollout',7,
       'The system supports professional judgement; it does not replace it.')
card(s,.71,2.2,5.82,3.99)
label(s,'Cross-report candidates (not approved)',.99,2.48,5.0,PURPLE)
text(s,'Application whitelisting gaps',.99,3.06,5.05,.33,18,NAVY,True)
text(s,'2018 + 2019  /  4 cited records  /  3 agency names',1.0,3.43,5.08,.44,12,SLATE)
shape(s,1.0,4.04,4.98,.017,BORDER)
text(s,'Patch-management weaknesses',.99,4.36,5.15,.33,18,NAVY,True)
text(s,'2018 + 2019  /  6 cited records  /  5 agency names',1.0,4.74,5.08,.44,12,SLATE)
text(s,'Both require source-level auditor confirmation before being described as recurring findings.',1.0,5.54,5.13,.46,11,SLATE)
card(s,6.75,2.2,5.86,3.99)
label(s,'Development path',7.04,2.48,4.89,TEAL)
steps=[
 ('1','Benchmark','Create blinded, expert-labelled pages; measure missed weaknesses and duplicates.'),
 ('2','Harden','Require agency-specific evidence; improve atomicity and deduplication.'),
 ('3','Pilot','Add reviewer sign-off, data governance and access controls before real audit use.')]
for i,(n,hd,body) in enumerate(steps):
    y=2.97+i*.99
    shape(s,7.05,y,.34,.34,'E6F7F3',None,True)
    text(s,n,7.12,y+.06,.2,.2,11,TEAL,True)
    text(s,hd,7.53,y,4.44,.31,16,NAVY,True)
    text(s,body,7.53,y+.35,4.60,.51,12,SLATE)
text(s,'AI USE DISCLOSURE  ChatGPT/Codex assisted with design, implementation, source-first evaluation and slide drafting. Findings remain provisional; no professional audit sign-off.',.84,6.36,11.71,.61,11,SLATE)

out=ROOT/'outputs'/'GovTech_Audit_Assurance_Proposal.pptx'
out.parent.mkdir(exist_ok=True)
prs.save(out)
with zipfile.ZipFile(out) as z:
    assert z.testzip() is None
print('saved',out,'bytes',out.stat().st_size,'slides',len(prs.slides))

